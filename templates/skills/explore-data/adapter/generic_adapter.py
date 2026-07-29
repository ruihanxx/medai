from __future__ import annotations

import argparse
import bz2
import csv
import gzip
import io
import json
import lzma
import mimetypes
import os
import sqlite3
import tarfile
import wave
import zipfile
from collections import OrderedDict
from pathlib import Path
from typing import Any, BinaryIO
from urllib.parse import quote
from xml.etree import ElementTree

REFERENCE_FILES = [
    "general_scientific_formats.md",
    "bioinformatics_genomics_formats.md",
    "microscopy_imaging_formats.md",
    "proteomics_metabolomics_formats.md",
]
REFERENCE_ROOT = Path(__file__).parents[1] / "references"
COMPRESSION_SUFFIXES = {
    ".gz": "gzip",
    ".gzip": "gzip",
    ".bz2": "bzip2",
    ".xz": "xz",
}
DANGEROUS_FORMATS = {
    "ckpt",
    "dill",
    "joblib",
    "pickle",
    "pkl",
    "pt",
    "pth",
}
TEXT_FORMATS = {
    "bed",
    "bedgraph",
    "csv",
    "fa",
    "fasta",
    "fastq",
    "fna",
    "fq",
    "gff",
    "gff3",
    "gtf",
    "ini",
    "json",
    "jsonl",
    "log",
    "md",
    "sam",
    "tab",
    "toml",
    "tsv",
    "txt",
    "vcf",
    "xml",
    "yaml",
    "yml",
}
IMAGE_FORMATS = {
    "bmp",
    "gif",
    "jpeg",
    "jpg",
    "ome.tif",
    "ome.tiff",
    "png",
    "tif",
    "tiff",
}


class DatasetAdapter:
    def __init__(self, raw_root: str | Path):
        self.raw_root = Path(raw_root).expanduser().resolve()
        if not self.raw_root.is_dir():
            raise ValueError(f"Raw data directory does not exist: {self.raw_root}")

    def list_resources(self, max_files: int = 100_000) -> list[dict[str, Any]]:
        return self._scan(max_files)["catalog"]

    def sample(
        self,
        relative_path: str,
        limit: int = 5,
        max_field_chars: int = 512,
        max_sample_bytes: int = 1024 * 1024,
    ) -> dict[str, Any]:
        if limit < 0:
            raise ValueError("limit must be non-negative")
        if max_field_chars < 1 or max_sample_bytes < 1:
            raise ValueError("Sample bounds must be positive")

        path = self._resolve(relative_path)
        if not path.is_file():
            raise FileNotFoundError(path)
        file_format, compression = self._format(path.name)
        result: dict[str, Any] = {
            "path": Path(relative_path).as_posix(),
            "format": file_format,
            "compression": compression,
            "size_bytes": path.stat().st_size,
            "reference_candidates": self._reference_candidates(file_format),
            "structure": {},
            "sample": None,
            "error": None,
        }

        try:
            if file_format in DANGEROUS_FORMATS:
                result["structure"] = {
                    "safe_to_load": False,
                    "mime_type": mimetypes.guess_type(path.name)[0],
                }
                result["error"] = (
                    "Unsafe serialized format was not loaded because deserialization "
                    "may execute code"
                )
                return result
            if file_format in {"zip", "tar"}:
                self._sample_archive(path, file_format, limit, result)
                return result
            if file_format in {"sqlite", "sqlite3", "db"}:
                self._sample_sqlite(path, limit, max_field_chars, result)
                return result
            if file_format == "wav":
                self._sample_wave(path, limit, result)
                return result
            if file_format in {"npy", "npz"}:
                self._sample_numpy(path, file_format, limit, max_field_chars, result)
                return result
            if file_format in {"parquet", "feather", "arrow"}:
                self._sample_arrow(path, file_format, limit, max_field_chars, result)
                return result
            if file_format in {"xlsx", "xlsm"}:
                self._sample_excel(path, limit, max_field_chars, result)
                return result
            if file_format in {"h5", "hdf5", "h5ad"}:
                self._sample_hdf5(path, limit, result)
                return result
            if file_format in IMAGE_FORMATS:
                self._sample_image(path, limit, result)
                return result
            if file_format == "dcm":
                self._sample_dicom(path, limit, max_field_chars, result)
                return result
            if file_format in {"nii", "nii.gz"}:
                self._sample_nifti(path, result)
                return result
            if file_format in {"bam", "cram"}:
                self._sample_alignment(path, file_format, limit, max_field_chars, result)
                return result

            prefix, source_truncated = self._read_prefix(
                path,
                compression,
                max_sample_bytes,
            )
            if file_format in {"csv", "tsv", "tab"}:
                delimiter = "," if file_format == "csv" else "\t"
                self._sample_delimited(
                    prefix,
                    delimiter,
                    limit,
                    max_field_chars,
                    source_truncated,
                    result,
                )
                return result
            if file_format == "jsonl":
                self._sample_jsonl(
                    prefix,
                    limit,
                    max_field_chars,
                    source_truncated,
                    result,
                )
                return result
            if file_format == "json":
                self._sample_json(
                    prefix,
                    limit,
                    max_field_chars,
                    source_truncated,
                    result,
                )
                return result
            if file_format == "xml":
                self._sample_xml(prefix, limit, source_truncated, result)
                return result
            if file_format in {"fa", "fasta", "fna"}:
                self._sample_fasta(
                    prefix,
                    limit,
                    max_field_chars,
                    source_truncated,
                    result,
                )
                return result
            if file_format in {"fastq", "fq"}:
                self._sample_fastq(
                    prefix,
                    limit,
                    max_field_chars,
                    source_truncated,
                    result,
                )
                return result
            if file_format in TEXT_FORMATS or self._looks_like_text(prefix):
                self._sample_text(
                    prefix,
                    limit,
                    max_field_chars,
                    source_truncated,
                    result,
                )
                return result

            result["structure"] = {
                "kind": "binary",
                "mime_type": mimetypes.guess_type(path.name)[0],
                "prefix_bytes_read": len(prefix),
                "printable_ratio": self._printable_ratio(prefix),
            }
            result["sample"] = {
                "prefix_hex": prefix[:64].hex(),
                "source_truncated": source_truncated,
            }
            result["error"] = "No safe format-specific reader was available"
        except ImportError as exc:
            result["error"] = f"Missing optional reader dependency: {exc}"
        except Exception as exc:  # noqa: BLE001 - preserve per-file exploration failures.
            result["error"] = f"{exc.__class__.__name__}: {exc}"
        return result

    def build_inventory(
        self,
        selected_files: list[str] | None = None,
        rows: int = 5,
        max_field_chars: int = 512,
        max_files: int = 100_000,
        max_sampled_files: int = 20,
        max_sample_bytes: int = 1024 * 1024,
    ) -> dict[str, Any]:
        scan = self._scan(max_files)
        files_to_sample = (
            selected_files
            if selected_files is not None
            else scan["representatives"][:max_sampled_files]
        )
        explored_files = [
            self.sample(
                relative_path,
                limit=rows,
                max_field_chars=max_field_chars,
                max_sample_bytes=max_sample_bytes,
            )
            for relative_path in files_to_sample
        ]
        warnings = list(scan["warnings"])
        if scan["truncated"]:
            warnings.append(
                f"Directory scan stopped after {max_files} files; rescan relevant "
                "subdirectories before relying on coverage"
            )
        for item in explored_files:
            if item["error"]:
                warnings.append(f"{item['path']}: {item['error']}")

        dataset_id = self.raw_root.parent.name if self.raw_root.name == "raw" else self.raw_root.name
        return {
            "schema_version": 1,
            "dataset": {
                "id": dataset_id,
                "root": str(self.raw_root),
                "adapter": "generic",
                "status": "explored",
            },
            "scan": {
                "files_scanned": scan["files_scanned"],
                "bytes_scanned": scan["bytes_scanned"],
                "truncated": scan["truncated"],
                "limits": {
                    "max_files": max_files,
                    "max_sampled_files": max_sampled_files,
                    "rows_per_file": rows,
                    "max_field_chars": max_field_chars,
                    "max_sample_bytes": max_sample_bytes,
                },
            },
            "catalog": scan["catalog"],
            "explored_files": explored_files,
            "warnings": warnings,
        }

    def write_inventory(
        self,
        output: str | Path,
        selected_files: list[str] | None = None,
        rows: int = 5,
        max_field_chars: int = 512,
        max_files: int = 100_000,
        max_sampled_files: int = 20,
        max_sample_bytes: int = 1024 * 1024,
    ) -> dict[str, Any]:
        destination = Path(output).expanduser().resolve()
        if destination.is_relative_to(self.raw_root):
            raise ValueError("Inventory output must not be inside the raw data directory")
        payload = self.build_inventory(
            selected_files=selected_files,
            rows=rows,
            max_field_chars=max_field_chars,
            max_files=max_files,
            max_sampled_files=max_sampled_files,
            max_sample_bytes=max_sample_bytes,
        )
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_text(
            json.dumps(payload, indent=2, ensure_ascii=False) + "\n",
            encoding="utf-8",
        )
        return payload

    def _scan(self, max_files: int) -> dict[str, Any]:
        if max_files < 1:
            raise ValueError("max_files must be positive")
        groups: OrderedDict[tuple[str, str, str | None], dict[str, Any]] = OrderedDict()
        representatives: list[str] = []
        warnings: list[str] = []
        files_scanned = 0
        bytes_scanned = 0
        truncated = False

        for current_root, directory_names, file_names in os.walk(
            self.raw_root,
            followlinks=False,
        ):
            current = Path(current_root)
            kept_directories: list[str] = []
            for name in sorted(directory_names):
                directory = current / name
                if directory.is_symlink():
                    warnings.append(
                        f"Skipped symbolic-link directory: "
                        f"{directory.relative_to(self.raw_root).as_posix()}"
                    )
                else:
                    kept_directories.append(name)
            directory_names[:] = kept_directories

            for name in sorted(file_names):
                if files_scanned >= max_files:
                    truncated = True
                    break
                path = current / name
                relative_path = path.relative_to(self.raw_root).as_posix()
                if path.is_symlink():
                    warnings.append(f"Skipped symbolic-link file: {relative_path}")
                    continue
                try:
                    size_bytes = path.stat().st_size
                except OSError as exc:
                    warnings.append(f"Could not stat {relative_path}: {exc}")
                    continue

                file_format, compression = self._format(name)
                directory = path.parent.relative_to(self.raw_root).as_posix()
                key = (directory, file_format, compression)
                if key not in groups:
                    groups[key] = {
                        "directory": directory,
                        "format": file_format,
                        "compression": compression,
                        "file_count": 0,
                        "total_bytes": 0,
                        "examples": [],
                    }
                    representatives.append(relative_path)
                group = groups[key]
                group["file_count"] += 1
                group["total_bytes"] += size_bytes
                if len(group["examples"]) < 3:
                    group["examples"].append(relative_path)
                files_scanned += 1
                bytes_scanned += size_bytes
            if truncated:
                break

        return {
            "catalog": list(groups.values()),
            "representatives": representatives,
            "files_scanned": files_scanned,
            "bytes_scanned": bytes_scanned,
            "truncated": truncated,
            "warnings": warnings,
        }

    def _sample_archive(
        self,
        path: Path,
        file_format: str,
        limit: int,
        result: dict[str, Any],
    ) -> None:
        members: list[dict[str, Any]] = []
        has_more = False
        if file_format == "zip":
            with zipfile.ZipFile(path) as archive:
                for index, item in enumerate(archive.infolist()):
                    if index >= limit:
                        has_more = True
                        break
                    members.append(
                        {
                            "name": item.filename,
                            "size_bytes": item.file_size,
                            "compressed_bytes": item.compress_size,
                            "is_directory": item.is_dir(),
                        }
                    )
        else:
            with tarfile.open(path, mode="r:*") as archive:
                for index, item in enumerate(archive):
                    if index >= limit:
                        has_more = True
                        break
                    members.append(
                        {
                            "name": item.name,
                            "size_bytes": item.size,
                            "is_directory": item.isdir(),
                        }
                    )
        result["structure"] = {"kind": "archive", "extracted": False}
        result["sample"] = {"members": members, "members_truncated": has_more}

    def _sample_sqlite(
        self,
        path: Path,
        limit: int,
        max_field_chars: int,
        result: dict[str, Any],
    ) -> None:
        uri = f"file:{quote(str(path))}?mode=ro"
        connection = sqlite3.connect(uri, uri=True)
        connection.row_factory = sqlite3.Row
        try:
            connection.execute("PRAGMA query_only = ON")
            table_names = [
                row[0]
                for row in connection.execute(
                    "SELECT name FROM sqlite_master "
                    "WHERE type = 'table' ORDER BY name LIMIT 20"
                )
            ]
            tables: list[dict[str, Any]] = []
            for table_name in table_names:
                quoted_name = table_name.replace('"', '""')
                columns = [
                    row[1]
                    for row in connection.execute(
                        f'PRAGMA table_info("{quoted_name}")'
                    )
                ]
                rows = [
                    {
                        key: self._json_safe(row[key], max_field_chars)
                        for key in row.keys()
                    }
                    for row in connection.execute(
                        f'SELECT * FROM "{quoted_name}" LIMIT ?',
                        (limit,),
                    )
                ]
                tables.append(
                    {
                        "name": table_name,
                        "columns": columns,
                        "sample_rows": rows,
                    }
                )
            result["structure"] = {"kind": "sqlite", "table_count_shown": len(tables)}
            result["sample"] = {"tables": tables}
        finally:
            connection.close()

    @staticmethod
    def _sample_wave(
        path: Path,
        limit: int,
        result: dict[str, Any],
    ) -> None:
        with wave.open(str(path), "rb") as audio:
            frame_bytes = audio.readframes(limit)
            result["structure"] = {
                "channels": audio.getnchannels(),
                "sample_width_bytes": audio.getsampwidth(),
                "sample_rate_hz": audio.getframerate(),
                "frame_count": audio.getnframes(),
                "compression": audio.getcomptype(),
            }
            result["sample"] = {"frame_prefix_hex": frame_bytes.hex()}

    def _sample_numpy(
        self,
        path: Path,
        file_format: str,
        limit: int,
        max_field_chars: int,
        result: dict[str, Any],
    ) -> None:
        import numpy

        if file_format == "npy":
            array = numpy.load(path, mmap_mode="r", allow_pickle=False)
            result["structure"] = {
                "shape": list(array.shape),
                "dtype": str(array.dtype),
            }
            result["sample"] = {
                "values": self._json_safe(
                    array.flat[:limit].tolist(),
                    max_field_chars,
                )
            }
            return

        arrays: list[dict[str, Any]] = []
        with numpy.load(path, allow_pickle=False) as archive:
            for name in archive.files[:20]:
                array = archive[name]
                arrays.append(
                    {
                        "name": name,
                        "shape": list(array.shape),
                        "dtype": str(array.dtype),
                        "values": self._json_safe(
                            array.flat[:limit].tolist(),
                            max_field_chars,
                        ),
                    }
                )
        result["structure"] = {"array_count_shown": len(arrays)}
        result["sample"] = {"arrays": arrays}

    def _sample_arrow(
        self,
        path: Path,
        file_format: str,
        limit: int,
        max_field_chars: int,
        result: dict[str, Any],
    ) -> None:
        if file_format == "parquet":
            import pyarrow.parquet as parquet

            source = parquet.ParquetFile(path)
            batch = next(source.iter_batches(batch_size=max(limit, 1)), None)
            rows = batch.to_pylist()[:limit] if batch is not None else []
            result["structure"] = {
                "columns": source.schema.names,
                "row_groups": source.num_row_groups,
            }
        else:
            import pyarrow
            import pyarrow.ipc as ipc

            with pyarrow.memory_map(str(path), "r") as source:
                reader = ipc.open_file(source)
                batch = reader.get_batch(0).slice(0, limit) if reader.num_record_batches else None
                rows = batch.to_pylist() if batch is not None else []
                result["structure"] = {
                    "columns": reader.schema.names,
                    "record_batches": reader.num_record_batches,
                }
        result["sample"] = {
            "rows": self._json_safe(rows, max_field_chars),
            "rows_returned": len(rows),
        }

    def _sample_excel(
        self,
        path: Path,
        limit: int,
        max_field_chars: int,
        result: dict[str, Any],
    ) -> None:
        import openpyxl

        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook[workbook.sheetnames[0]]
            rows = [
                [self._json_safe(value, max_field_chars) for value in row]
                for row in list(sheet.iter_rows(values_only=True, max_row=limit + 1))
            ]
            result["structure"] = {
                "sheet_names": workbook.sheetnames,
                "selected_sheet": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
            }
            result["sample"] = {"rows": rows}
        finally:
            workbook.close()

    @staticmethod
    def _sample_hdf5(
        path: Path,
        limit: int,
        result: dict[str, Any],
    ) -> None:
        import h5py

        items: list[dict[str, Any]] = []
        with h5py.File(path, "r") as handle:

            def collect(name: str, item: Any) -> bool | None:
                if len(items) >= limit:
                    return True
                entry: dict[str, Any] = {
                    "path": name,
                    "kind": "dataset" if isinstance(item, h5py.Dataset) else "group",
                }
                if isinstance(item, h5py.Dataset):
                    entry["shape"] = list(item.shape)
                    entry["dtype"] = str(item.dtype)
                items.append(entry)
                return None

            handle.visititems(collect)
        result["structure"] = {"kind": "hdf5"}
        result["sample"] = {"items": items, "items_truncated": len(items) >= limit}

    @staticmethod
    def _sample_image(
        path: Path,
        limit: int,
        result: dict[str, Any],
    ) -> None:
        from PIL import Image

        with Image.open(path) as image:
            image.seek(0)
            pixels = []
            for index, value in enumerate(image.getdata()):
                if index >= limit:
                    break
                pixels.append(value)
            result["structure"] = {
                "format": image.format,
                "mode": image.mode,
                "size": list(image.size),
                "frames": getattr(image, "n_frames", 1),
            }
            result["sample"] = {"first_frame_pixels": pixels}

    def _sample_dicom(
        self,
        path: Path,
        limit: int,
        max_field_chars: int,
        result: dict[str, Any],
    ) -> None:
        import pydicom

        dataset = pydicom.dcmread(path, stop_before_pixels=True)
        elements: list[dict[str, Any]] = []
        for index, element in enumerate(dataset):
            if index >= limit:
                break
            elements.append(
                {
                    "tag": str(element.tag),
                    "name": element.name,
                    "value": self._json_safe(element.value, max_field_chars),
                }
            )
        result["structure"] = {"kind": "dicom", "pixels_loaded": False}
        result["sample"] = {"metadata_elements": elements}

    @staticmethod
    def _sample_nifti(
        path: Path,
        result: dict[str, Any],
    ) -> None:
        import nibabel

        image = nibabel.load(path)
        result["structure"] = {
            "kind": "nifti",
            "shape": list(image.shape),
            "dtype": str(image.get_data_dtype()),
            "affine": image.affine.tolist(),
        }
        result["sample"] = {"voxel_values_loaded": False}

    def _sample_alignment(
        self,
        path: Path,
        file_format: str,
        limit: int,
        max_field_chars: int,
        result: dict[str, Any],
    ) -> None:
        import pysam

        mode = "rb" if file_format == "bam" else "rc"
        with pysam.AlignmentFile(path, mode) as handle:
            records = []
            for index, record in enumerate(handle.fetch(until_eof=True)):
                if index >= limit:
                    break
                records.append(self._json_safe(record.to_string(), max_field_chars))
            result["structure"] = {
                "kind": file_format,
                "references": list(handle.references),
                "lengths": list(handle.lengths),
            }
            result["sample"] = {"records": records}

    def _sample_delimited(
        self,
        prefix: bytes,
        delimiter: str,
        limit: int,
        max_field_chars: int,
        source_truncated: bool,
        result: dict[str, Any],
    ) -> None:
        text = self._complete_text_prefix(prefix, source_truncated)
        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
        columns = reader.fieldnames or []
        rows: list[dict[str, str]] = []
        fields_truncated = False
        for index, row in enumerate(reader):
            if index >= limit:
                break
            sampled_row: dict[str, str] = {}
            for column in columns:
                value, truncated = self._truncate(row.get(column), max_field_chars)
                sampled_row[column] = value
                fields_truncated = fields_truncated or truncated
            rows.append(sampled_row)
        result["structure"] = {"columns": columns}
        result["sample"] = {
            "rows": rows,
            "rows_returned": len(rows),
            "fields_truncated": fields_truncated,
            "source_truncated": source_truncated,
        }

    def _sample_jsonl(
        self,
        prefix: bytes,
        limit: int,
        max_field_chars: int,
        source_truncated: bool,
        result: dict[str, Any],
    ) -> None:
        text = self._complete_text_prefix(prefix, source_truncated)
        values = []
        for line in text.splitlines()[:limit]:
            if line.strip():
                values.append(self._json_safe(json.loads(line), max_field_chars))
        result["structure"] = {"kind": "json-lines"}
        result["sample"] = {
            "values": values,
            "values_returned": len(values),
            "source_truncated": source_truncated,
        }

    def _sample_json(
        self,
        prefix: bytes,
        limit: int,
        max_field_chars: int,
        source_truncated: bool,
        result: dict[str, Any],
    ) -> None:
        text = prefix.decode("utf-8", errors="replace")
        if source_truncated:
            value, fields_truncated = self._truncate(text, max_field_chars)
            result["structure"] = {"kind": "json", "parsed": False}
            result["sample"] = {
                "prefix": value,
                "source_truncated": True,
                "fields_truncated": fields_truncated,
            }
            result["error"] = "JSON exceeded the bounded read and was not partially parsed"
            return
        data = json.loads(text)
        if isinstance(data, list):
            sampled = data[:limit]
        elif isinstance(data, dict):
            sampled = dict(list(data.items())[:limit])
        else:
            sampled = data
        result["structure"] = {
            "kind": "json",
            "top_level_type": type(data).__name__,
            "length": len(data) if isinstance(data, (list, dict)) else None,
        }
        result["sample"] = {"value": self._json_safe(sampled, max_field_chars)}

    @staticmethod
    def _sample_xml(
        prefix: bytes,
        limit: int,
        source_truncated: bool,
        result: dict[str, Any],
    ) -> None:
        if source_truncated:
            result["structure"] = {"kind": "xml", "parsed": False}
            result["sample"] = {
                "prefix": prefix[:512].decode("utf-8", errors="replace"),
                "source_truncated": True,
            }
            result["error"] = "XML exceeded the bounded read and was not partially parsed"
            return
        root = ElementTree.fromstring(prefix)
        tags = []
        for index, element in enumerate(root.iter()):
            if index >= limit:
                break
            tags.append(element.tag)
        result["structure"] = {"kind": "xml", "root_tag": root.tag}
        result["sample"] = {"element_tags": tags}

    def _sample_fasta(
        self,
        prefix: bytes,
        limit: int,
        max_field_chars: int,
        source_truncated: bool,
        result: dict[str, Any],
    ) -> None:
        text = self._complete_text_prefix(prefix, source_truncated)
        records: list[dict[str, Any]] = []
        current_id: str | None = None
        sequence: list[str] = []
        for line in text.splitlines():
            if line.startswith(">"):
                if current_id is not None:
                    records.append(
                        self._sequence_record(current_id, sequence, max_field_chars)
                    )
                    if len(records) >= limit:
                        break
                current_id = line[1:]
                sequence = []
            elif current_id is not None:
                sequence.append(line.strip())
        if current_id is not None and len(records) < limit:
            records.append(self._sequence_record(current_id, sequence, max_field_chars))
        result["structure"] = {"kind": "fasta"}
        result["sample"] = {
            "records": records,
            "source_truncated": source_truncated,
        }

    def _sample_fastq(
        self,
        prefix: bytes,
        limit: int,
        max_field_chars: int,
        source_truncated: bool,
        result: dict[str, Any],
    ) -> None:
        lines = self._complete_text_prefix(prefix, source_truncated).splitlines()
        records = []
        for offset in range(0, min(len(lines) - 3, limit * 4), 4):
            identifier, sequence, separator, quality = lines[offset : offset + 4]
            if not identifier.startswith("@") or not separator.startswith("+"):
                break
            records.append(
                {
                    "id": self._truncate(identifier[1:], max_field_chars)[0],
                    "sequence_length": len(sequence),
                    "sequence_prefix": self._truncate(sequence, max_field_chars)[0],
                    "quality_prefix": self._truncate(quality, max_field_chars)[0],
                }
            )
        result["structure"] = {"kind": "fastq"}
        result["sample"] = {
            "records": records,
            "source_truncated": source_truncated,
        }

    def _sample_text(
        self,
        prefix: bytes,
        limit: int,
        max_field_chars: int,
        source_truncated: bool,
        result: dict[str, Any],
    ) -> None:
        text = self._complete_text_prefix(prefix, source_truncated)
        lines = []
        fields_truncated = False
        for line in text.splitlines()[:limit]:
            value, truncated = self._truncate(line, max_field_chars)
            lines.append(value)
            fields_truncated = fields_truncated or truncated
        result["structure"] = {
            "kind": "text",
            "encoding": "utf-8-or-replacement",
        }
        result["sample"] = {
            "lines": lines,
            "lines_returned": len(lines),
            "fields_truncated": fields_truncated,
            "source_truncated": source_truncated,
        }

    def _reference_candidates(self, file_format: str) -> list[dict[str, str]]:
        candidates: list[dict[str, str]] = []
        needle = f".{file_format}".casefold()
        for filename in REFERENCE_FILES:
            path = REFERENCE_ROOT / filename
            with path.open(encoding="utf-8") as handle:
                for line in handle:
                    heading = line.strip()
                    if heading.startswith("### ") and needle in heading.casefold():
                        candidates.append(
                            {
                                "reference": f"references/{filename}",
                                "section": heading.removeprefix("### "),
                            }
                        )
        return candidates

    def _resolve(self, relative_path: str) -> Path:
        candidate = Path(relative_path)
        if candidate.is_absolute():
            raise ValueError("Dataset file paths must be relative to the raw root")
        resolved = (self.raw_root / candidate).resolve()
        if not resolved.is_relative_to(self.raw_root):
            raise ValueError(f"Dataset file escapes the raw root: {relative_path}")
        return resolved

    @staticmethod
    def _format(filename: str) -> tuple[str, str | None]:
        lower_name = filename.casefold()
        for compound in ("ome.tiff", "ome.tif"):
            if lower_name.endswith(f".{compound}"):
                return compound, None
        if lower_name.endswith(".tar.gz") or lower_name.endswith(".tgz"):
            return "tar", "gzip"
        if lower_name.endswith(".tar.bz2"):
            return "tar", "bzip2"
        if lower_name.endswith(".tar.xz"):
            return "tar", "xz"
        if lower_name.endswith(".nii.gz"):
            return "nii.gz", "gzip"

        suffixes = Path(lower_name).suffixes
        if not suffixes:
            return "unknown", None
        compression = COMPRESSION_SUFFIXES.get(suffixes[-1])
        if compression and len(suffixes) >= 2:
            return suffixes[-2].lstrip("."), compression
        return suffixes[-1].lstrip("."), compression

    @staticmethod
    def _read_prefix(
        path: Path,
        compression: str | None,
        max_bytes: int,
    ) -> tuple[bytes, bool]:
        handle: BinaryIO
        if compression == "gzip":
            handle = gzip.open(path, "rb")
        elif compression == "bzip2":
            handle = bz2.open(path, "rb")
        elif compression == "xz":
            handle = lzma.open(path, "rb")
        else:
            handle = path.open("rb")
        with handle:
            content = handle.read(max_bytes + 1)
        return content[:max_bytes], len(content) > max_bytes

    @staticmethod
    def _complete_text_prefix(prefix: bytes, source_truncated: bool) -> str:
        text = prefix.decode("utf-8", errors="replace")
        if source_truncated and "\n" in text:
            return text.rsplit("\n", 1)[0]
        return text

    @staticmethod
    def _looks_like_text(prefix: bytes) -> bool:
        if not prefix:
            return True
        return DatasetAdapter._printable_ratio(prefix) >= 0.85

    @staticmethod
    def _printable_ratio(prefix: bytes) -> float:
        if not prefix:
            return 1.0
        printable = sum(
            byte in {9, 10, 13} or 32 <= byte <= 126
            for byte in prefix
        )
        return printable / len(prefix)

    @staticmethod
    def _truncate(value: Any, max_chars: int) -> tuple[str, bool]:
        text = "" if value is None else str(value)
        if len(text) <= max_chars:
            return text, False
        return text[:max_chars], True

    @classmethod
    def _json_safe(cls, value: Any, max_chars: int) -> Any:
        if value is None or isinstance(value, (bool, int, float)):
            return value
        if isinstance(value, bytes):
            return value[:max_chars].hex()
        if isinstance(value, str):
            return cls._truncate(value, max_chars)[0]
        if isinstance(value, dict):
            return {
                cls._truncate(key, max_chars)[0]: cls._json_safe(item, max_chars)
                for key, item in list(value.items())[:100]
            }
        if isinstance(value, (list, tuple)):
            return [cls._json_safe(item, max_chars) for item in value[:100]]
        return cls._truncate(value, max_chars)[0]

    @classmethod
    def _sequence_record(
        cls,
        identifier: str,
        sequence_parts: list[str],
        max_field_chars: int,
    ) -> dict[str, Any]:
        sequence = "".join(sequence_parts)
        return {
            "id": cls._truncate(identifier, max_field_chars)[0],
            "sequence_length": len(sequence),
            "sequence_prefix": cls._truncate(sequence, max_field_chars)[0],
        }


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Create a bounded, read-only inventory for an arbitrary dataset."
    )
    parser.add_argument("raw_root", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--file", action="append", default=None, dest="selected_files")
    parser.add_argument("--rows", type=int, default=5)
    parser.add_argument("--max-field-chars", type=int, default=512)
    parser.add_argument("--max-files", type=int, default=100_000)
    parser.add_argument("--max-sampled-files", type=int, default=20)
    parser.add_argument("--max-sample-bytes", type=int, default=1024 * 1024)
    args = parser.parse_args()

    DatasetAdapter(args.raw_root).write_inventory(
        args.output,
        selected_files=args.selected_files,
        rows=args.rows,
        max_field_chars=args.max_field_chars,
        max_files=args.max_files,
        max_sampled_files=args.max_sampled_files,
        max_sample_bytes=args.max_sample_bytes,
    )


if __name__ == "__main__":
    main()
